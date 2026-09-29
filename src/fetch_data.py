"""
fetch_data.py
下載油價與美債研究所需資料，對齊至共同期間後輸出 Excel。

資料來源
- FRED (St. Louis Fed)：美債殖利率、油價、股市（景氣控制變數）
- 紐約聯儲 ACM 模型：期限溢價與風險中立殖利率（預期政策利率成分）

使用方式
    export FRED_API_KEY=你的金鑰        # 或寫在專案根目錄的 .env
    python src/fetch_data.py
    python src/fetch_data.py --strict-end   # 結束日也取「所有序列都有資料」的最後一天

每次執行都重新下載完整歷史並覆寫檔案，因此可安全地重複排程執行。
"""

from __future__ import annotations

import argparse
import io
import os
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------------
# 設定
# ----------------------------------------------------------------------------
ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "data" / "oil_bond_data.xlsx"

FRED_URL = "https://api.stlouisfed.org/fred/series/observations"

# FRED 序列：代碼 -> (欄位說明, 單位)
FRED_SERIES = {
    "DGS2": ("2年期美債殖利率 (CMT)", "%"),
    "DGS10": ("10年期美債殖利率 (CMT)", "%"),
    "DGS30": ("30年期美債殖利率 (CMT)", "%"),
    "DCOILBRENTEU": ("Brent 原油現貨價格", "USD/bbl"),
    "DCOILWTICO": ("WTI 原油現貨價格", "USD/bbl"),
    "NASDAQCOM": ("NASDAQ 綜合指數（景氣控制變數）", "index"),
    # TIPS 相關序列（約 2003 年起），用於 breakeven 拆解
    "DFII10": ("10年期 TIPS 實質殖利率", "%"),
    "T10YIE": ("10年期通膨平衡率 (DGS10 - DFII10)", "%"),
    "T5YIFR": ("5y5y 遠期通膨預期", "%"),
}

# 起始較晚的補充序列：不參與共同期間起點的判斷，起點之前顯示空白。
# 否則整個資料集會被推到 2003 年，失去 1990 年的衝擊期間。
LATE_START_SERIES = {"DFII10", "T10YIE", "T5YIFR"}

# ACM 期限溢價檔案（紐約聯儲）。若網址變更，只需修改此處。
ACM_URL = (
    "https://www.newyorkfed.org/medialibrary/media/research/"
    "data_indicators/ACMTermPremium.xls"
)
# ACMY = 模型配適殖利率；ACMTP = 期限溢價；ACMRNY = 風險中立殖利率（預期短率路徑）
ACM_SERIES = {
    "ACMY02": ("ACM 2年期配適殖利率", "%"),
    "ACMTP02": ("ACM 2年期期限溢價", "%"),
    "ACMRNY02": ("ACM 2年期風險中立殖利率", "%"),
    "ACMY10": ("ACM 10年期配適殖利率", "%"),
    "ACMTP10": ("ACM 10年期期限溢價", "%"),
    "ACMRNY10": ("ACM 10年期風險中立殖利率", "%"),
}

MAX_RETRIES = 4
TIMEOUT = 60


# ----------------------------------------------------------------------------
# 工具函式
# ----------------------------------------------------------------------------
def load_api_key() -> str:
    """優先讀環境變數，其次讀專案根目錄的 .env。金鑰不寫進程式碼。"""
    key = os.environ.get("FRED_API_KEY")
    env_file = ROOT / ".env"
    if not key and env_file.exists():
        for line in env_file.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line.startswith("FRED_API_KEY="):
                key = line.split("=", 1)[1].strip().strip('"').strip("'")
    if not key:
        sys.exit("找不到 FRED_API_KEY：請設定環境變數或在 .env 中加入 FRED_API_KEY=...")
    return key


def get_with_retry(url: str, **kwargs) -> requests.Response:
    """帶指數退避的 GET，處理 FRED 偶發的 429 / 5xx。"""
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = requests.get(url, timeout=TIMEOUT, **kwargs)
            if resp.status_code == 200:
                return resp
            if resp.status_code in (429, 500, 502, 503, 504):
                raise requests.HTTPError(f"HTTP {resp.status_code}")
            resp.raise_for_status()
        except requests.RequestException as err:
            if attempt == MAX_RETRIES:
                raise
            wait = 2 ** attempt
            print(f"  重試 {attempt}/{MAX_RETRIES - 1}（{err}），{wait} 秒後再試")
            time.sleep(wait)
    raise RuntimeError("unreachable")


# ----------------------------------------------------------------------------
# 下載
# ----------------------------------------------------------------------------
def fetch_fred_series(series_id: str, api_key: str) -> pd.Series:
    params = {
        "series_id": series_id,
        "api_key": api_key,
        "file_type": "json",
        "observation_start": "1900-01-01",
    }
    data = get_with_retry(FRED_URL, params=params).json()
    obs = pd.DataFrame(data["observations"])
    # FRED 以 "." 表示缺值（假日等）
    values = pd.to_numeric(obs["value"].replace(".", pd.NA), errors="coerce")
    s = pd.Series(values.values, index=pd.to_datetime(obs["date"]), name=series_id)
    return s.dropna()


def fetch_acm() -> pd.DataFrame:
    resp = get_with_retry(ACM_URL, headers={"User-Agent": "Mozilla/5.0"})
    sheets = pd.read_excel(io.BytesIO(resp.content), sheet_name=None)

    # 找日頻工作表（名稱含 Daily），找不到就退而使用第一張
    name = next((n for n in sheets if "daily" in n.lower()), next(iter(sheets)))
    df = sheets[name]
    df.columns = [str(c).strip().upper() for c in df.columns]

    date_col = next(c for c in df.columns if "DATE" in c)
    df[date_col] = pd.to_datetime(df[date_col], errors="coerce")
    df = df.dropna(subset=[date_col]).set_index(date_col).sort_index()
    df.index.name = None

    missing = [c for c in ACM_SERIES if c not in df.columns]
    if missing:
        raise KeyError(f"ACM 檔案（工作表 {name}）缺少欄位：{missing}")
    return df[list(ACM_SERIES)].apply(pd.to_numeric, errors="coerce")


# ----------------------------------------------------------------------------
# 對齊與中繼資料
# ----------------------------------------------------------------------------
def longest_gap(s: pd.Series) -> tuple[str, int]:
    """回傳期間內最長連續缺值的區段與天數（營業日），用於提醒資料斷層。"""
    isna = s.isna()
    if not isna.any():
        return "", 0
    groups = (isna != isna.shift()).cumsum()
    runs = s[isna].groupby(groups[isna])
    longest = max(runs, key=lambda kv: len(kv[1]))[1]
    return f"{longest.index[0].date()} ~ {longest.index[-1].date()}", len(longest)


def build_dataset(raw: dict[str, pd.Series], strict_end: bool) -> tuple[pd.DataFrame, pd.Timestamp, pd.Timestamp]:
    firsts = {k: v.first_valid_index() for k, v in raw.items()}
    lasts = {k: v.last_valid_index() for k, v in raw.items()}

    # 共同起點只看核心序列；補充序列（LATE_START_SERIES）在各自起點之前留白
    start = max(v for k, v in firsts.items() if k not in LATE_START_SERIES)
    end = min(lasts.values()) if strict_end else max(lasts.values())

    idx = pd.bdate_range(start, end)
    df = pd.DataFrame({k: v.reindex(idx) for k, v in raw.items()}, index=idx)
    # 移除全部欄位皆缺值的日子（例如所有市場都休市）
    df = df.dropna(how="all")
    df.index.name = "date"
    return df, start, end


def build_meta(raw: dict[str, pd.Series], df: pd.DataFrame, start, end, strict_end: bool) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows = []
    for code, s in raw.items():
        desc, unit = {**FRED_SERIES, **ACM_SERIES}[code]
        source = "FRED" if code in FRED_SERIES else "NY Fed ACM"
        # 斷層只從該序列自己的起始日開始算，補充序列起點前的空白不算斷層
        gap_range, gap_len = longest_gap(df[code].loc[s.first_valid_index():])
        rows.append({
            "代碼": code,
            "說明": desc,
            "單位": unit,
            "來源": source,
            "決定共同起點": "否（補充序列）" if code in LATE_START_SERIES else "是",
            "原始起始日": s.first_valid_index().date(),
            "原始最新日": s.last_valid_index().date(),
            "期間內有效筆數": int(df[code].notna().sum()),
            "期間內最長缺值區段": gap_range,
            "缺值營業日數": gap_len,
        })
    meta = pd.DataFrame(rows)

    run_time = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    notes = pd.DataFrame({
        "項目": ["更新時間", "共同期間", "結束日規則", "補充序列", "缺值說明"],
        "內容": [
            run_time,
            f"{start.date()} ~ {end.date()}",
            "所有序列共同的最後一天" if strict_end else "任一序列的最新資料日（較晚更新的序列在近期會顯示空白）",
            f"{', '.join(sorted(LATE_START_SERIES))} 起始較晚，不決定共同起點，起點前為空白；僅用於 2003 年後的 breakeven 拆解。",
            "空白代表該日無資料（假日或序列停編），未做任何插補；請於分析階段自行處理。",
        ],
    })
    return meta, notes


# ----------------------------------------------------------------------------
# 輸出 Excel
# ----------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", start_color="1F3864")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial")


def style_sheet(ws, number_format: str | None = None, date_cols: tuple[int, ...] = ()):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            if cell.column in date_cols:
                cell.number_format = "yyyy-mm-dd"
            elif number_format and isinstance(cell.value, (int, float)):
                cell.number_format = number_format
    ws.freeze_panes = "B2"
    for i, col in enumerate(ws.columns, start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200])
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, width * 1.3), 60)


def write_excel(df: pd.DataFrame, meta: pd.DataFrame, notes: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")  # 先寫暫存檔，成功後再取代，避免排程中斷留下損毀檔
    with pd.ExcelWriter(tmp, engine="openpyxl") as writer:
        df.reset_index().to_excel(writer, sheet_name="data", index=False)
        meta.to_excel(writer, sheet_name="series_info", index=False)
        notes.to_excel(writer, sheet_name="notes", index=False)

        style_sheet(writer.sheets["data"], number_format="0.00", date_cols=(1,))
        style_sheet(writer.sheets["series_info"], date_cols=(6, 7))
        style_sheet(writer.sheets["notes"])
    tmp.replace(path)


# ----------------------------------------------------------------------------
# 主程式
# ----------------------------------------------------------------------------
def main() -> None:
    parser = argparse.ArgumentParser(description="下載油價與美債研究資料")
    parser.add_argument("--strict-end", action="store_true",
                        help="結束日取所有序列共同的最後一天（預設為最新可得日）")
    parser.add_argument("--output", type=Path, default=OUTPUT_PATH)
    args = parser.parse_args()

    api_key = load_api_key()
    raw: dict[str, pd.Series] = {}

    for code in FRED_SERIES:
        print(f"下載 FRED {code} ...")
        raw[code] = fetch_fred_series(code, api_key)

    print("下載 ACM 期限溢價 ...")
    acm = fetch_acm()
    for code in ACM_SERIES:
        raw[code] = acm[code].dropna()

    df, start, end = build_dataset(raw, args.strict_end)
    meta, notes = build_meta(raw, df, start, end, args.strict_end)
    write_excel(df, meta, notes, args.output)

    print(f"\n完成：{args.output}")
    print(f"共同期間 {start.date()} ~ {end.date()}，共 {len(df):,} 筆")
    flagged = meta[meta["缺值營業日數"] > 20]
    if not flagged.empty:
        print("注意：以下序列在期間內有較長斷層")
        print(flagged[["代碼", "期間內最長缺值區段", "缺值營業日數"]].to_string(index=False))


if __name__ == "__main__":
    main()
